#!/usr/bin/env python3
from flask import Flask, render_template, request, jsonify, send_file
from flask_cors import CORS
from werkzeug.utils import secure_filename
import pandas as pd
import openpyxl
import json
import os
from urllib.parse import quote

app = Flask(__name__, static_folder='.', static_url_path='')
CORS(app, resources={r"/api/*": {"origins": "*"}})
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024
app.config['UPLOAD_FOLDER'] = 'uploads'

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Store widget file mappings
widget_files = {}

WIDGET_CONFIG = {
    1: {"name": "Travel NA Bookings FY'27 YTD", "file": "04.Travel NA Bookings.xlsx"},
    2: {"name": "Travel NA Bookings Monthly", "file": "06.Travel NA Bookings Current Month.xlsx"},
    3: {"name": "Travel NA Pipeline Deals", "file": "05. Travel NA Pipeline Deals.xlsx"},
    4: {"name": "Travel NA Client Connect Monthly", "file": "Travel NA Client Connect-Current Month.xlsx"},
    5: {"name": "Travel NA Meeting Data Monthly", "file": "Travel NA Meeting Data Monthly.xlsx"},
    6: {"name": "Travel NA Open Deals", "file": "Travel NA Open Deals.xlsx"},
    7: {"name": "Budget Vs BE by QTR", "file": "04.Travel NA BE FY27_Finance.xlsx"},
    8: {"name": "Open Positions (RRs)", "file": "KPI_Report_Travel NA.xlsx"},
    9: {"name": "Action Tracker", "file": "Travel NA_Action Tracker.xlsx"}
}

def find_data_table(filepath):
    """Find the actual data table in Excel by looking for numeric columns"""
    wb = openpyxl.load_workbook(filepath)
    
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        
        # Scan rows to find a header row with multiple non-empty cells
        for row_idx in range(1, min(100, ws.max_row + 1)):
            row_cells = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))[0]
            non_empty = [c for c in row_cells if c is not None and str(c).strip()]
            
            # If we find a row with 3+ non-empty cells, check if next rows have numeric data
            if len(non_empty) >= 3:
                next_row = list(ws.iter_rows(min_row=row_idx+1, max_row=row_idx+1, values_only=True))[0]
                numeric_cells = sum(1 for c in next_row if c is not None and isinstance(c, (int, float)))
                
                # If next row has numeric data, this is likely the header
                if numeric_cells >= 2:
                    return sheet, row_idx
    
    return None, None

def parse_workbook(filepath):
    """Return the dashboard payload for an Excel workbook."""
    sheet_name, header_row = find_data_table(filepath)
    data = {}
    xls = pd.ExcelFile(filepath)

    for sheet in xls.sheet_names:
        try:
            if sheet == sheet_name and header_row:
                df = pd.read_excel(filepath, sheet_name=sheet, header=header_row-1, nrows=50)
            else:
                df = pd.read_excel(filepath, sheet_name=sheet, nrows=20)

            if len(df) > 0 and len(df.columns) > 1:
                data[sheet] = {
                    'columns': df.columns.tolist(),
                    'rows': df.where(pd.notna(df), '').astype(str).values.tolist()[:15],
                    'row_count': len(df),
                    'header_row': header_row if sheet == sheet_name else None
                }
        except Exception as e:
            data[sheet] = {'error': str(e)}

    return {
        'filename': os.path.basename(filepath),
        'sheets': xls.sheet_names,
        'data': data,
        'detected_header_row': header_row,
        'detected_sheet': sheet_name
    }

def configured_file(widget_id):
    """Find the uploaded copy for a widget without relying on process memory."""
    prefix = f'widget_{widget_id}_'
    configured_name = secure_filename(WIDGET_CONFIG[widget_id]['file'])
    direct_path = os.path.join(app.config['UPLOAD_FOLDER'], configured_name)
    if os.path.isfile(direct_path):
        return direct_path

    candidates = [
        os.path.join(app.config['UPLOAD_FOLDER'], name)
        for name in os.listdir(app.config['UPLOAD_FOLDER'])
        if name.startswith(prefix) and os.path.isfile(os.path.join(app.config['UPLOAD_FOLDER'], name))
    ]
    if not candidates:
        return None

    exact = [path for path in candidates if os.path.basename(path) == f'{prefix}{configured_name}']
    return exact[0] if exact else max(candidates, key=os.path.getmtime)

@app.route('/')
def index():
    return app.send_static_file('index.html')

@app.route('/api/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    
    file = request.files['file']
    widget_id = request.form.get('widget_id', type=int)
    
    if not file or file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    if not widget_id or widget_id < 1 or widget_id > 9:
        return jsonify({'error': 'Invalid widget ID'}), 400
    
    try:
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], f'widget_{widget_id}_{filename}')
        file.save(filepath)
        
        # Store the mapping for download
        widget_files[widget_id] = filepath
        print(f"Stored widget {widget_id} file: {filepath}")
        
        result = parse_workbook(filepath)
        result.update({'success': True, 'widget_id': widget_id,
                       'filepath': filepath,
                       'file_download_url': f'/api/download/{widget_id}'})
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': f'Failed to parse file: {str(e)}'}), 400

@app.route('/api/widget/<int:widget_id>/data')
def widget_data(widget_id):
    """Load the configured workbook on every dashboard refresh."""
    if widget_id not in WIDGET_CONFIG:
        return jsonify({'error': 'Invalid widget ID'}), 400

    filepath = configured_file(widget_id)
    if not filepath:
        return jsonify({
            'error': f'No Excel source found for widget {widget_id}. '
                     'Place the configured workbook in the uploads folder.'
        }), 404

    try:
        result = parse_workbook(filepath)
        result.update({'success': True, 'widget_id': widget_id,
                       'file_download_url': f'/api/download/{widget_id}'})
        return jsonify(result)
    except Exception as e:
        return jsonify({'error': f'Failed to parse file: {str(e)}'}), 400

@app.route('/api/download/<int:widget_id>')
def download_file(widget_id):
    """Serve uploaded file for download"""
    if widget_id < 1 or widget_id > 9:
        return jsonify({'error': 'Invalid widget ID'}), 400
    
    filepath = widget_files.get(widget_id) or configured_file(widget_id)
    if not filepath:
        return jsonify({'error': 'File not found in the uploads folder.'}), 404
    
    if not os.path.exists(filepath):
        print(f"File path does not exist: {filepath}")
        return jsonify({'error': f'File not found at path: {filepath}'}), 404
    
    print(f"Serving file: {filepath}")
    return send_file(filepath, as_attachment=True, download_name=os.path.basename(filepath))

@app.route('/api/widget-config')
def get_config():
    return jsonify(WIDGET_CONFIG)

if __name__ == '__main__':
    app.run(host='localhost', port=5000, debug=False)
